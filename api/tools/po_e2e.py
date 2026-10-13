"""End-to-end check of Purchase Order generation (PO BRD, 11 acceptance criteria) through the real API
(in-process; real Gemini for the AI-assisted draft and the uploaded-PO extraction).

    cd api && python tools/po_e2e.py            # throwaway data dirs
    cd api && python tools/po_e2e.py --keep     # write into the real data dirs (for a browser check)
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
from datetime import date, timedelta
from pathlib import Path

if "--keep" not in sys.argv:
    tmp = Path(tempfile.mkdtemp(prefix="cx-po-e2e-"))
    os.environ["RFQ_DIR"], os.environ["QUOTE_DIR"] = str(tmp / "rfqs"), str(tmp / "quotes")
    os.environ["PO_DIR"], os.environ["BUYER_DIR"] = str(tmp / "pos"), str(tmp / "buyers")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.stdout.reconfigure(encoding="utf-8")

import docx  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app import main, store  # noqa: E402

c = TestClient(main.app)
BUYER = {"X-Role": "buyer", "X-User-Id": "voltedge-power", "X-User-Name": "VoltEdge Power Pvt Ltd"}
seller = lambda sid: {"X-Role": "seller", "X-User-Id": sid}  # noqa: E731
fails = []


def check(name, cond, detail=""):
    print(("PASS " if cond else "FAIL ") + name + (f"  [{detail}]" if detail and not cond else ""))
    if not cond:
        fails.append(name)


def ok(r):
    assert r.status_code == 200, f"{r.request.method} {r.request.url} -> {r.status_code} {r.text[:400]}"
    return r.json()


def errors(p):
    return [x for x in p["checks"] if x["level"] == "error"]


# ── setup: buyer profile, transformer RFQ, two template quotations ───────────
prof = ok(c.get("/api/buyer/profile", headers=BUYER))
check("profile starts with the standard terms library", len(prof["terms"]) >= 10)
ok(c.put("/api/buyer/profile", json={"name": "VoltEdge Power Pvt Ltd", "registered_office": "Plot 4, MIDC Chakan, Pune",
                                     "gstin": "27AAACV1234F1Z5", "billing_address": "Accounts, Plot 4, MIDC Chakan, Pune",
                                     "shipping_addresses": ["Plant 1, MIDC Chakan, Pune"], "signatory_name": "Anita Rao",
                                     "signatory_designation": "Head of Procurement"}, headers=BUYER))
rid = ok(c.post("/api/rfqs/from-template", json={"template_id": "transformer"}, headers=BUYER))["id"]
ok(c.put(f"/api/rfqs/{rid}/slots", json={"values": {"buyer": "VoltEdge Power Pvt Ltd", "quantity": 50,
                                                    "standard": "IS 2026"}}, headers=BUYER))
ok(c.post(f"/api/rfqs/{rid}/approve", headers=BUYER))
ok(c.post(f"/api/rfqs/{rid}/confirm", headers=BUYER))
gate = ok(c.get(f"/api/rfqs/{rid}", headers=BUYER))["gate"]
sellers = [s for s in ("ref-remi", "ref-transcore") if gate.get(s, {}).get("passed")]
assert len(sellers) == 2, gate
deadline = (date.today() + timedelta(days=7)).isoformat()
sent = ok(c.post(f"/api/rfqs/{rid}/quote-requests", json={"seller_ids": sellers, "deadline": deadline}, headers=BUYER))
req_of = {q["seller_id"]: q["id"] for q in sent["requests"]}
for sid, price in (("ref-remi", 410000), ("ref-transcore", 395000)):
    d = ok(c.post(f"/api/seller/requests/{req_of[sid]}/draft", json={"mode": "template"}, headers=seller(sid)))["draft"]
    d.update({"lines": [{**li, "unit_price": price} for li in d["lines"]], "gst_pct": 18, "freight": 25000,
              "lead_time_days": 45, "validity_days": 90, "payment_terms": "60 days from delivery", "payment_days": 60,
              "compliance": d["compliance"]})
    ok(c.put(f"/api/seller/requests/{req_of[sid]}/quotation", json={"quotation": d}, headers=seller(sid)))
    ok(c.post(f"/api/seller/requests/{req_of[sid]}/submit", headers=seller(sid)))
remi, tc = req_of["ref-remi"], req_of["ref-transcore"]
quote = ok(c.get(f"/api/rfqs/{rid}/quotes", headers=BUYER))
remi_q = next(r["quotation"] for r in quote["requests"] if r["id"] == remi)


def outcomes():
    return {r["seller_id"]: r["outcome"] for r in ok(c.get(f"/api/rfqs/{rid}/quotes", headers=BUYER))["requests"]}


def supplier_sees():
    return {s: (ok(c.get(f"/api/seller/requests/{req_of[s]}", headers=seller(s)))["outcome"],
                len(ok(c.get("/api/seller/pos", headers=seller(s))))) for s in sellers}


# ── C1 Accept opens a PO draft, nothing is accepted ──────────────────────────
p = ok(c.post(f"/api/rfqs/{rid}/quotes/{remi}/select", headers=BUYER))
check("C1 Accept opens a PO draft", p["status"] == "draft" and p["request_id"] == remi)
oc = outcomes()
check("C1 quotation shows Selected (PO pending), not Accepted", oc == {"ref-remi": "selected", "ref-transcore": None}, oc)
check("C6 supplier sees nothing before issue; other supplier unchanged",
      supplier_sees() == {"ref-remi": (None, 0), "ref-transcore": (None, 0)}, supplier_sees())
x = c.post(f"/api/rfqs/{rid}/quotes/{tc}/select", headers=BUYER)
check("C7 a second quote cannot be selected while one is pending", x.status_code == 409, x.status_code)

# ── C2 template: pre-filled, items / qty / price locked ──────────────────────
p = ok(c.post(f"/api/rfqs/{rid}/po/start", json={"mode": "template"}, headers=BUYER))
f = p["fields"]
check("C2 template PO pre-filled from quotation + RFQ + profile",
      f["supplier_name"] == "Remi" and f["buyer_name"] == "VoltEdge Power Pvt Ltd" and f["ship_to"]
      and f["payment_terms"] == "60 days from delivery" and p["totals"]["total"] == remi_q["totals"]["landed_total"])
p = ok(c.put(f"/api/rfqs/{rid}/po", json={"lines": [{"desc": "Tampered", "qty": 1, "unit": "u", "unit_price": 1}],
                                          "gst_pct": 0, "freight": 0}, headers=BUYER))
check("C2 item, quantity and price cannot be edited (server re-copies the quotation)",
      p["lines"][0]["unit_price"] == 410000 and p["lines"][0]["qty"] == remi_q["lines"][0]["qty"] and p["gst_pct"] == 18)

# ── C3 required fields + expired quotation block issue ───────────────────────
codes = {(x["code"], x["message"]) for x in errors(p)}
check("C3 missing required field flagged (delivery date)", any("delivery date" in m for _, m in codes), codes)
check("C3 placeholder in the standard Disputes clause flagged", any(k == "placeholder" for k, _ in codes), codes)
x = c.post(f"/api/rfqs/{rid}/po/issue", headers=BUYER)
check("C3 issue refused while a required field is empty", x.status_code == 400, x.status_code)
terms = [t for t in p["terms"] if t["id"] != "t13"]
p = ok(c.put(f"/api/rfqs/{rid}/po", json={"fields": {"delivery_date": (date.today() + timedelta(days=60)).isoformat(),
                                                     "inspection": "Routine tests witnessed before dispatch"},
                                          "terms": terms}, headers=BUYER))
check("C3 checks pass once completed", not errors(p), errors(p))
rq = store.load_request(remi)
real = rq["submitted_at"]
rq["submitted_at"] = (date.today() - timedelta(days=120)).isoformat() + "T10:00:00+00:00"
store.save_request(rq)
p2 = ok(c.get(f"/api/rfqs/{rid}/po", headers=BUYER))
x = c.post(f"/api/rfqs/{rid}/po/issue", headers=BUYER)
check("C3 expired quotation blocks issue", any(e["code"] == "quote_expired" for e in errors(p2)) and x.status_code == 400,
      (errors(p2), x.status_code))
rq["submitted_at"] = real
store.save_request(rq)

# saved PO template
t = ok(c.post(f"/api/rfqs/{rid}/po/save-template", json={"name": "Transformer PO"}, headers=BUYER))["templates"]
p = ok(c.post(f"/api/rfqs/{rid}/po/start", json={"mode": "template", "template_id": t[0]["id"]}, headers=BUYER))
check("C8 saved PO template reused (open fields + clauses kept)",
      p["fields"]["inspection"] == "Routine tests witnessed before dispatch" and [x["id"] for x in p["terms"]] == [x["id"] for x in terms])

# ── C8/C9 AI-assisted draft ───────────────────────────────────────────────────
p = ok(c.post(f"/api/rfqs/{rid}/po/start", json={"mode": "ai"}, headers=BUYER))
qids = [q["id"] for q in p["questions"]]
print("   AI questions:", qids, "| terms chosen:", [x["id"] for x in p["terms"]], "| schedule:", p["schedule"])
check("C9 AI asks only what it cannot fill (no supplier / item / price / payment / known address questions)",
      set(qids) <= {"delivery_date", "inspection", "special_terms"} and "delivery_date" in qids, qids)
check("C2 AI draft keeps the quotation's items and prices", p["lines"] == ok(c.get(f"/api/rfqs/{rid}/po", headers=BUYER))["lines"]
      and p["totals"]["total"] == remi_q["totals"]["landed_total"])
lib = {x["id"] for x in prof["terms"]}
check("C9 AI only selects saved clauses, never writes new ones", all(x["id"] in lib for x in p["terms"]))
check("AI schedule has no echoed header row", all(r[0] != "Stage" for r in p["schedule"]), p["schedule"])
p = ok(c.post(f"/api/rfqs/{rid}/po/ai-answers", json={"answers": {
    "delivery_date": (date.today() + timedelta(days=50)).isoformat(), "inspection": "Pre-dispatch inspection at works"}},
    headers=BUYER))
if any(e["code"] == "placeholder" for e in errors(p)):
    p = ok(c.put(f"/api/rfqs/{rid}/po", json={"terms": [x for x in p["terms"] if x["id"] != "t13"]}, headers=BUYER))
check("C11 AI draft passes the same pre-issue checks", not errors(p), errors(p))
pv = ok(c.post(f"/api/rfqs/{rid}/po/preview", json={}, headers=BUYER))["blocks"]
check("preview renders the PO document", pv[0]["text"] == "Purchase Order" and any(b["t"] == "h1" for b in pv))

# ── C7 cancel selection, choose another quote ────────────────────────────────
ok(c.delete(f"/api/rfqs/{rid}/po", headers=BUYER))
check("C7 cancel: draft discarded, quotation back to Submitted", outcomes() == {"ref-remi": None, "ref-transcore": None}
      and c.get(f"/api/rfqs/{rid}/po", headers=BUYER).status_code == 404)
ok(c.post(f"/api/rfqs/{rid}/quotes/{tc}/select", headers=BUYER))
check("C7 another quote can then be selected", outcomes()["ref-transcore"] == "selected")
ok(c.delete(f"/api/rfqs/{rid}/po", headers=BUYER))
ok(c.post(f"/api/rfqs/{rid}/quotes/{remi}/select", headers=BUYER))


# ── C10 upload: mismatches flagged, then corrected ───────────────────────────
def po_file(supplier, qty, price, total):
    d = docx.Document()
    d.add_heading("PURCHASE ORDER", 0)
    for line in (f"PO No: VE-PO-2026-0099    PO Date: {date.today().isoformat()}",
                 "Buyer: VoltEdge Power Pvt Ltd, Plot 4, MIDC Chakan, Pune. GSTIN 27AAACV1234F1Z5",
                 "Bill to: Accounts, Plot 4, MIDC Chakan, Pune", "Ship to: Plant 1, MIDC Chakan, Pune",
                 f"Supplier: {supplier}, Bhopal", f"Required delivery date: {(date.today() + timedelta(days=55)).isoformat()}",
                 "Payment: 60 days from delivery", "Authorised signatory: Anita Rao, Head of Procurement"):
        d.add_paragraph(line)
    tb = d.add_table(rows=2, cols=4)
    for j, v in enumerate(("Description", "Qty", "Unit rate (Rs)", "Amount (Rs)")):
        tb.cell(0, j).text = v
    for j, v in enumerate(("Distribution transformer 100 kVA", str(qty), f"{price:,}", f"{qty * price:,}")):
        tb.cell(1, j).text = v
    d.add_paragraph("GST 18% extra. Freight Rs 25,000.")
    d.add_paragraph(f"Total PO value incl. GST and freight: Rs {total:,.0f}")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


qty = remi_q["lines"][0]["qty"]
bad = po_file("KCS Engineering", qty + 5, 400000, (qty + 5) * 400000 * 1.18 + 25000)
p = ok(c.post(f"/api/rfqs/{rid}/po/upload", files={"file": ("our_po.docx", bad, "application/octet-stream")}, headers=BUYER))
mm = [e["message"] for e in errors(p) if e["code"] == "mismatch"]
print("   upload mismatches:", mm)
check("C10 uploaded PO: supplier, quantity, price and total mismatches all flagged",
      any("Supplier" in m for m in mm) and any("quantity" in m for m in mm) and any("unit price" in m for m in mm)
      and any("total" in m.lower() for m in mm), mm)
x = c.post(f"/api/rfqs/{rid}/po/issue", headers=BUYER)
check("C10 mismatching upload cannot be issued", x.status_code == 400, x.status_code)
good = po_file("Remi", qty, 410000, remi_q["totals"]["landed_total"])
p = ok(c.post(f"/api/rfqs/{rid}/po/upload", files={"file": ("our_po.docx", good, "application/octet-stream")}, headers=BUYER))
print("   upload extracted:", {k: p["fields"].get(k) for k in ("po_number", "supplier_name", "delivery_date", "ship_to")},
      [(li["qty"], li["unit_price"]) for li in p["lines"]], p["gst_pct"], p["freight"], p.get("stated_total"))
mm = [e["message"] for e in errors(p) if e["code"] == "mismatch"]
check("C10 matching upload has no mismatch", not mm, mm)
check("C11 upload goes through the same checks", not errors(p), errors(p))

# ── C5 delivery failure: nothing accepted ────────────────────────────────────
real_deliver = main.deliver_po
main.deliver_po = lambda p: (_ for _ in ()).throw(RuntimeError("supplier inbox unavailable"))
x = c.post(f"/api/rfqs/{rid}/po/issue", headers=BUYER)
main.deliver_po = real_deliver
check("C5 delivery failure -> quotation not Accepted, PO still draft, supplier sees nothing",
      x.status_code == 502 and outcomes()["ref-remi"] == "selected"
      and ok(c.get(f"/api/rfqs/{rid}/po", headers=BUYER))["status"] == "draft"
      and supplier_sees() == {"ref-remi": (None, 0), "ref-transcore": (None, 0)}, (x.status_code, outcomes()))

# ── C4/C6 Generate PO ─────────────────────────────────────────────────────────
r = ok(c.post(f"/api/rfqs/{rid}/po/issue", headers=BUYER))
check("C4 Generate PO: Accepted + Not selected + PO issued, in one action",
      r["po"]["status"] == "issued" and outcomes() == {"ref-remi": "accepted", "ref-transcore": "not_selected"}, outcomes())
lst = ok(c.get("/api/seller/pos", headers=seller("ref-remi")))
check("C6 supplier now sees the PO with a notification (unread)", len(lst) == 1 and lst[0]["unread"])
v = ok(c.get(f"/api/seller/pos/{lst[0]['id']}", headers=seller("ref-remi")))
check("C6 opening the PO clears the notification", not ok(c.get("/api/seller/pos", headers=seller("ref-remi")))[0]["unread"]
      and v["blocks"][0]["text"] == "Purchase Order")
x = c.get(f"/api/seller/pos/{lst[0]['id']}/file?role=seller&uid=ref-remi")
check("upload: the buyer's own file is delivered to the supplier", x.status_code == 200 and x.content == good)
x = c.get(f"/api/seller/pos/{lst[0]['id']}/po.docx?role=seller&uid=ref-remi")
check("supplier downloads the PO .docx", x.status_code == 200 and x.content[:2] == b"PK")
check("other supplier sees no PO, only its Not selected status",
      supplier_sees()["ref-transcore"] == ("not_selected", 0)
      and c.get(f"/api/seller/pos/{lst[0]['id']}", headers=seller("ref-transcore")).status_code == 404)
x = c.post(f"/api/rfqs/{rid}/quotes/{tc}/select", headers=BUYER)
check("after issue no other quote can be selected", x.status_code == 409, x.status_code)
sv = ok(c.get(f"/api/seller/requests/{remi}", headers=seller("ref-remi")))
check("winning supplier's request links to the PO", sv["outcome"] == "accepted" and sv["po_id"] == lst[0]["id"])

print(f"\n{'ALL PASS' if not fails else f'{len(fails)} FAILED: ' + '; '.join(fails)}  (RFQ {rid})")
print(json.dumps({"rfq": rid}, indent=0)) if "--keep" in sys.argv else None
sys.exit(1 if fails else 0)
