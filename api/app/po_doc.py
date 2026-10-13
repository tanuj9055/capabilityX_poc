"""Purchase order document: preview blocks + .docx, in the layout of the buyer's PO templates
(Purchase_Order_Template_*.docx): parties table, order value, key controls, then numbered sections.
The same blocks feed the live preview (web components/Preview.jsx) and the .docx the supplier receives.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

BASE = Path(__file__).resolve().parent / "data" / "templates" / "po_base.docx"
TODO = "[To be completed?]"  # preview only: highlighted for the buyer
NAVY, ZEBRA, BORDER, TEAL, INK = "17365D", "F7F9FB", "D9E0E7", "159A95", "222222"


def inr(v) -> str:
    """Indian digit grouping: 21890500 -> 2,18,90,500."""
    if v is None:
        return "-"
    neg, v = v < 0, abs(float(v))
    whole, frac = divmod(round(v, 2), 1)
    s = str(int(whole))
    if len(s) > 3:
        head, tail = s[:-3], s[-3:]
        head = ",".join([head[max(0, i - 2):i] for i in range(len(head), 0, -2)][::-1])
        s = f"{head},{tail}"
    if frac:
        s += f"{frac:.2f}"[1:]
    return ("-" if neg else "") + s


def _num(v) -> str:
    return "-" if v is None else (f"{v:,.0f}" if float(v).is_integer() else f"{v:,.2f}")


def _long(iso: str | None) -> str:
    try:
        d = datetime.fromisoformat(str(iso)[:10]).date()
    except (TypeError, ValueError):
        return str(iso or "")
    return f"{d.day} {d.strftime('%B %Y')}"


def blocks(po: dict, draft: bool = False) -> list[dict]:
    f, t, rfq = po["fields"], po["totals"], po["rfq"]
    gap = TODO if draft else ""
    v = lambda k: f.get(k) or gap  # noqa: E731
    out: list[dict] = []

    def p(text, kind="body"):
        out.append({"t": kind, "text": text})

    def table(rows):
        out.append({"t": "table", "header": True, "rows": rows})

    p("Purchase Order", "title")
    p(rfq["title"], "sub")
    p(f"{f.get('po_number') or 'PO'}  |  {rfq['ref']}  |  {f.get('quote_ref') or ''}".upper(), "meta")
    table([["Buyer field", "Buyer or PO information", "Supplier field", "Supplier information"],
           ["Buyer", v("buyer_name"), "Supplier", v("supplier_name")],
           ["Registered office", v("buyer_address"), "Supplier address", f.get("supplier_address") or ""],
           ["Buyer GSTIN", f.get("buyer_gstin") or "", "Supplier GSTIN", f.get("supplier_gstin") or ""],
           ["PO number", v("po_number"), "Supplier code", f.get("supplier_code") or ""],
           ["PO date", _long(f.get("po_date")) or gap, "Supplier contact", f.get("supplier_contact") or ""],
           ["RFQ reference", rfq["ref"], "Quotation reference", f.get("quote_ref") or ""],
           ["PO type", f.get("po_type") or "", "Currency", f.get("currency") or "INR"],
           ["Bill to", v("bill_to"), "Ship to", v("ship_to")],
           ["Delivery basis", f.get("delivery_basis") or "", "Payment", v("payment_terms")]])
    if f.get("intro"):
        p(f["intro"])
    if po.get("mode") == "upload" and po.get("file"):
        p(f"This PO is issued as the Buyer's attached file “{po['file']['name']}”. The confirmed order data below "
          "has been checked against the selected quotation.")
    p("Order value", "h2")
    table([["Order component", "Value INR", "Basis"],
           ["Order value before GST", inr(t["subtotal"]), "Firm, as per the commercial schedule"],
           [f"GST at {po.get('gst_pct') or 0:g} percent", inr(t["tax"]), "Against a valid tax invoice"]]
          + ([["Freight", inr(t["freight"]), "As quoted"]] if t.get("freight") else [])
          + [["Total PO value", inr(t["total"]), "Including GST" + (" and freight" if t.get("freight") else "")]])
    if po.get("key_controls"):
        p("Key controls", "h2")
        for x in po["key_controls"]:
            p(x, "bullet")

    p("1 Order Authorisation and Commercial Schedule", "h1")
    p("The Buyer issues this purchase order for the items below at the prices of the selected quotation. No charge "
      "outside this schedule is payable unless added through an authorised written PO amendment.")
    lump = lambda li: li.get("unit") == "lot" and li.get("qty") == 1  # noqa: E731
    table([["Line", "Description", "Quantity", "Unit price INR", "Value INR"]]
          + [[str(i), li["desc"], f"{_num(li.get('qty'))} {li.get('unit') or ''}".strip(),
              "-" if lump(li) else inr(li.get("unit_price")), inr(li.get("amount"))]
             for i, li in enumerate(po.get("lines") or [], 1)])
    p(f"PO value before GST: INR {inr(t['subtotal'])}", "bold")
    p(f"GST at {po.get('gst_pct') or 0:g} percent: INR {inr(t['tax'])}"
      + (f"; freight: INR {inr(t['freight'])}" if t.get("freight") else "") + f". Total PO value: INR {inr(t['total'])}.")

    p("2 Delivery Schedule", "h1")
    rows = po.get("schedule") or []
    if rows:
        table([["Stage", "Quantity", "Required timing", "Release and acceptance gate"]]
              + [[c or (gap if j == 2 else "") for j, c in enumerate(r)] for r in rows])
    p(f"Required delivery date: {_long(f.get('delivery_date')) or gap}", "bullet")
    p(f"Deliver to: {v('ship_to')}", "bullet")
    if f.get("delivery_basis"):
        p(f"Delivery basis: {f['delivery_basis']}", "bullet")
    p("Each delivery must carry the PO number, item, quantity, packing list and certificate of conformity.", "bullet")

    p("3 Accepted Clarifications and Buyer Conditions", "h1")
    p("Only the supplier assumptions expressly recorded below are accepted. All other quotation qualifications are "
      "superseded by this PO unless incorporated through an authorised amendment.")
    if po.get("clarifications"):
        table([["Quotation topic", "Buyer position and PO condition"]] + po["clarifications"])
    else:
        p("No supplier clarifications or deviations are accepted.")

    p("4 Inspection and Acceptance", "h1")
    for x in (f.get("inspection") or "").split("\n"):
        if x.strip():
            p(x.strip(), "bullet")
    p("Inspection, payment or use does not waive latent defects, warranty rights or the supplier's responsibility "
      "for conformity.")

    p("5 Tax Invoice and Payment", "h1")
    p(f"Payment term: {v('payment_terms')}. GST is payable at the legally applicable rate against a valid tax invoice.")

    p("6 Terms and Conditions", "h1")
    if po.get("terms"):
        for i, term in enumerate(po["terms"], 1):
            p(f"6.{i} {term['title']}", "h2")
            p(term["text"])
    else:
        p(TODO if draft else "No additional terms.")

    extra = [x.strip() for x in (f.get("special_terms") or "").split("\n") + (f.get("notes") or "").split("\n")
             if x.strip()]
    n = 7
    if extra:
        p("7 Special Terms and Notes", "h1")
        for x in extra:
            p(x, "bullet")
        n = 8
    p(f"{n} Document Precedence", "h1")
    p("If documents conflict, the following order applies unless a later authorised amendment expressly states otherwise.")
    table([["Priority", "Controlling document"], ["1", "Authorised PO amendments issued after this PO"],
           ["2", "This PO including its terms and conditions"], ["3", f"Buyer {rfq['ref']} and authorised addenda"],
           ["4", f"Supplier {f.get('quote_ref') or 'quotation'} only to the extent expressly accepted in this PO"]])
    p(f"{n + 1} Buyer Approval and Supplier Acknowledgement", "h1")
    who = ", ".join(x for x in (f.get("buyer_signatory"), f.get("buyer_designation")) if x) or gap
    table([["Buyer approval", "Name and designation", "Signature", "Date"],
           ["Authorised signatory", who, "", _long(f.get("po_date"))]])
    p(f"Issued by {f.get('buyer_name') or 'the Buyer'}. The Supplier acknowledges receipt and acceptance of this PO, "
      "subject only to corrections expressly agreed by the Buyer in writing.")
    table([["Supplier", "Authorised representative", "Signature and stamp", "Date"],
           [f.get("supplier_name") or "", "", "", ""]])
    return out


def render_docx(po: dict, path: Path) -> Path:
    d = docx.Document(str(BASE))
    body = d.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    f = po["fields"]
    _set_runs(d.sections[0].header.paragraphs, [f"{(f.get('buyer_name') or 'BUYER').upper()}  |  PURCHASE ORDER"])
    _set_runs(d.sections[0].footer.paragraphs,
              [f"{f.get('po_number') or 'PO'}  |  {po['rfq']['title'][:40]}  |  Page ", None, " of ", None])
    for b in blocks(po):
        k = b["t"]
        if k == "table":
            _table(d, b["rows"])
        elif k == "title":
            d.add_paragraph(b["text"], style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif k in ("sub", "meta"):
            para = d.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_after = Pt(5 if k == "sub" else 10)
            r = para.add_run(b["text"])
            r.bold = True
            r.font.name = "Aptos Display" if k == "sub" else "Aptos"
            r.font.size = Pt(13 if k == "sub" else 8)
            r.font.color.rgb = RGBColor.from_string(INK if k == "sub" else TEAL)
        elif k == "bold":
            d.add_paragraph(style="Body Text").add_run(b["text"]).bold = True
        else:
            d.add_paragraph(b["text"], style={"h1": "Heading 1", "h2": "Heading 2", "bullet": "List Bullet"}
                            .get(k, "Body Text"))
    d.save(str(path))
    return path


def _set_runs(paragraphs, texts: list[str | None]) -> None:
    """Replace the text runs of the first paragraph in order; None keeps a field run (PAGE / NUMPAGES)."""
    for para in paragraphs:
        runs = para.runs
        if not runs:
            continue
        for r, t in zip(runs, texts):
            if t is not None:
                r.text = t
        for r in runs[len(texts):]:
            r.text = ""
        return


def _table(d, rows: list[list[str]]) -> None:
    """Table styled like the PO templates: navy header, zebra rows, light borders, 8.5 pt Aptos."""
    ncols = max(len(r) for r in rows)
    tb = d.add_table(rows=len(rows), cols=ncols)
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    tb.autofit = False
    widths = _widths(d, rows, ncols)
    for gc, w in zip(tb._tbl.tblGrid.findall(qn("w:gridCol")), widths):  # LibreOffice reads the grid, Word the cells
        gc.set(qn("w:w"), str(int(w / 635)))  # EMU -> twips
    for i, row in enumerate(rows):
        tr = tb.rows[i]._tr
        trpr = tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            h = OxmlElement("w:tblHeader")
            h.set(qn("w:val"), "true")
            trpr.append(h)
        for j in range(ncols):
            cell = tb.cell(i, j)
            cell.width = widths[j]
            tcpr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), NAVY if i == 0 else (ZEBRA if i % 2 == 0 else "FFFFFF"))
            tcpr.append(shd)
            mar = OxmlElement("w:tcMar")
            for side, w in (("top", 90), ("start", 110), ("bottom", 90), ("end", 110)):
                e = OxmlElement(f"w:{side}")
                e.set(qn("w:w"), str(w))
                e.set(qn("w:type"), "dxa")
                mar.append(e)
            tcpr.append(mar)
            bd = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                e = OxmlElement(f"w:{side}")
                e.set(qn("w:val"), "single")
                e.set(qn("w:sz"), "6")
                e.set(qn("w:color"), BORDER)
                bd.append(e)
            tcpr.append(bd)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = para.add_run(row[j] if j < len(row) else "")
            r.font.name = "Aptos"
            r.font.size = Pt(8.5)
            r.bold = i == 0
            r.font.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else INK)
    d.add_paragraph().paragraph_format.space_after = Pt(0)


def _widths(d, rows, ncols):
    """Split the usable width by each column's longest text, keeping No./Step columns narrow."""
    sec = d.sections[0]
    usable = sec.page_width - sec.left_margin - sec.right_margin
    need = []
    for j in range(ncols):
        cells = [r[j] for r in rows if j < len(r)]
        word = max((len(w) for c in cells for w in c.split()), default=0)
        need.append(min(max(max(len(c) for c in cells), word + 6, 5), 70))
    total = sum(need)
    return [int(usable * n / total) for n in need]
